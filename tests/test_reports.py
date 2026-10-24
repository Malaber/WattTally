import io
from types import SimpleNamespace as NS

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

from app.reports import freeze_report, pdf_bytes, xlsx_bytes


def sample(policy="employer"):
    vehicle = NS(name="Tommy", plate="NW FS 97E", policy=policy)
    session = NS(
        id="session1",
        start="2026-10-01T10:00:00+00:00",
        end="2026-10-01T11:00:00+00:00",
        assignment="confirmed",
        reviewed_energy_kwh="3",
        payload={
            "energy_kwh": "3",
            "grid_kwh": "1",
            "solar_kwh": "2",
            "unknown_kwh": "0",
            "grid_cost": "0.3",
            "solar_cost": "0.16",
            "flags": [],
        },
    )
    tariff = NS(effective_from="2026-01-01T00:00:00+00:00", reimbursement_rate="0.30")
    return vehicle, session, tariff


def report(policy="employer"):
    v, s, t = sample(policy)
    snapshot, csv, digest = freeze_report(v, [s], "2026-10", "Europe/Berlin", [t])
    return NS(
        id="report1", policy=policy, month="2026-10", snapshot=snapshot, csv_text=csv, sha256=digest
    )


def test_employer_omits_all_solar_fields_in_all_formats():
    r = report()
    assert r.snapshot["amount_eur"] == "0.90"
    assert "solar" not in r.csv_text.lower()
    assert "grid" not in r.csv_text.lower()
    book = load_workbook(io.BytesIO(xlsx_bytes(r)))
    cells = " ".join(str(c.value) for sheet in book for row in sheet for c in row)
    assert "solar" not in cells.lower()
    assert "grid" not in cells.lower()
    text = " ".join(p.extract_text() for p in PdfReader(io.BytesIO(pdf_bytes(r))).pages)
    assert "solar" not in text.lower()
    assert "0.90" in text
    assert "NW FS 97E" in text


def test_actual_cost_uses_grid_and_solar_rates():
    r = report("actual")
    assert r.snapshot["amount_eur"] == "0.46"
    assert "Solar kWh" in r.csv_text


def test_suggested_assignment_cannot_be_exported():
    v, s, t = sample()
    s.assignment = "suggested"
    with pytest.raises(ValueError, match="No confirmed"):
        freeze_report(v, [s], "2026-10", "Europe/Berlin", [t])


def test_user_confirmed_meter_gaps_can_be_exported():
    v, s, t = sample()
    s.payload["flags"] = ["meter_reset"]
    assert freeze_report(v, [s], "2026-10", "Europe/Berlin", [t])[0]["amount_eur"] == "0.90"


def test_unknown_solar_allowed_only_for_employer():
    v, s, t = sample()
    s.payload["unknown_kwh"] = "3"
    s.payload["flags"] = ["missing_grid_or_power"]
    assert freeze_report(v, [s], "2026-10", "Europe/Berlin", [t])[0]["amount_eur"] == "0.90"
    v.policy = "actual"
    with pytest.raises(ValueError, match="without grid/solar"):
        freeze_report(v, [s], "2026-10", "Europe/Berlin", [t])


def test_csv_formula_injection_is_neutralized():
    v, s, t = sample()
    v.name = "=1+1"
    assert "'=1+1" in freeze_report(v, [s], "2026-10", "Europe/Berlin", [t])[1]


def test_dst_month_boundary_and_session_rounding():
    v, s, t = sample()
    s.start = "2026-09-30T22:30:00+00:00"
    s.end = "2026-10-01T00:00:00+00:00"
    assert freeze_report(v, [s], "2026-10", "Europe/Berlin", [t])[0]["session_count"] == 1


def test_small_allocation_gap_is_estimated_and_disclosed_in_all_formats():
    v, s, t = sample("actual")
    s.payload.update(energy_kwh="3.02", unknown_kwh="0.02", flags=["meter_gap"])
    s.reviewed_energy_kwh = "3.02"
    snapshot, csv, digest = freeze_report(v, [s], "2026-10", "Europe/Berlin", [t])
    assert snapshot["energy_kwh"] == "3.020"
    assert "Estimated split and cost for 0.020 kWh" in csv
    r = NS(id="r", policy="actual", month="2026-10", snapshot=snapshot, csv_text=csv, sha256=digest)
    book = load_workbook(io.BytesIO(xlsx_bytes(r)))
    assert "Estimated split" in book.active.cell(2, 13).value
    text = " ".join(page.extract_text() for page in PdfReader(io.BytesIO(pdf_bytes(r))).pages)
    assert "Estimated split" in text


def test_changed_total_requires_new_confirmation():
    v, s, t = sample()
    s.payload["energy_kwh"] = "4"
    with pytest.raises(ValueError, match="No confirmed"):
        freeze_report(v, [s], "2026-10", "Europe/Berlin", [t])


def test_zero_cycles_and_unreviewed_charges_do_not_block_confirmed_records():
    from copy import deepcopy

    v, s, t = sample()
    zero = deepcopy(s)
    zero.payload.update(energy_kwh="0", flags=["missing_meter_baseline"])
    zero.assignment = "unassigned"
    pending = deepcopy(s)
    pending.assignment = "suggested"
    snapshot, _, _ = freeze_report(v, [zero, pending, s], "2026-10", "Europe/Berlin", [t])
    assert snapshot["session_count"] == 1


def test_missing_tariff_gives_actionable_instruction():
    v, s, t = sample("actual")
    s.payload["flags"] = ["missing_tariff"]
    with pytest.raises(ValueError, match="Add a home electricity tariff"):
        freeze_report(v, [s], "2026-10", "Europe/Berlin", [t])


@pytest.mark.parametrize("energy,unknown", [("100", "0.101"), ("1", "0.02")])
def test_allocation_tolerance_has_absolute_and_relative_limits(energy, unknown):
    v, s, t = sample("actual")
    s.payload.update(energy_kwh=energy, unknown_kwh=unknown)
    s.reviewed_energy_kwh = energy
    with pytest.raises(ValueError, match="without grid/solar"):
        freeze_report(v, [s], "2026-10", "Europe/Berlin", [t])


def test_charge_is_counted_only_in_completion_month_in_report_timezone():
    v, s, t = sample()
    s.start = "2026-09-30T20:30:00+00:00"
    s.end = "2026-09-30T22:15:00+00:00"  # October 1 in Berlin.
    with pytest.raises(ValueError, match="No confirmed"):
        freeze_report(v, [s], "2026-09", "Europe/Berlin", [t])
    snapshot, _, _ = freeze_report(v, [s], "2026-10", "Europe/Berlin", [t])
    assert snapshot["session_count"] == 1
    assert "completion date" in snapshot["period_rule"]
    s.end = None
    with pytest.raises(ValueError, match="No confirmed"):
        freeze_report(v, [s], "2026-10", "Europe/Berlin", [t])


@pytest.mark.parametrize("count", [15, 70])
def test_company_pdf_across_page_breaks(count):
    import runpy
    from pathlib import Path

    example_report = runpy.run_path(
        str(Path(__file__).resolve().parents[1] / "scripts/generate_example_report.py")
    )["example_report"]

    report = example_report(count)
    pages = PdfReader(io.BytesIO(pdf_bytes(report))).pages
    assert len(pages) >= (2 if count == 15 else 3)
    text = "\n".join(page.extract_text() for page in pages)
    assert "Wallbox export" in text
    assert text.count("DEMO AB 123E") == 1
    assert text.count("Long Range") == 1
    assert "completion date" in text
    assert "Round each" not in text and "Whole sessions" not in text
    assert "solar" not in text.lower() and "disclaimer" not in text.lower()
    for index, page in enumerate(pages, 1):
        content = page.extract_text()
        assert report.sha256 in content
        assert "Energy kWh" in content
        assert "Amount EUR" in content
        assert float(page.mediabox.width) == pytest.approx(595.276, abs=0.01)
        assert float(page.mediabox.height) == pytest.approx(841.890, abs=0.01)
        assert str(index) in content


def test_company_details_are_frozen_in_canonical_csv():
    v, s, t = sample()
    v.vehicle_model = "Example E20"
    v.driver_name = "Alex Example"
    v.driver_address = "Example Street 12\n12345 Example Town"
    v.vin = "DISCOVERYVIN"
    v.report_vin = "REPORTVIN"
    snapshot, csv_text, digest = freeze_report(v, [s], "2026-10", "Europe/Berlin", [t])
    frozen = NS(
        id="details",
        policy="employer",
        month="2026-10",
        snapshot=snapshot,
        csv_text=csv_text,
        sha256=digest,
    )
    v.driver_name = "Changed Driver"
    assert freeze_report(v, [s], "2026-10", "Europe/Berlin", [t])[2] != digest
    text = " ".join(p.extract_text() for p in PdfReader(io.BytesIO(pdf_bytes(frozen))).pages)
    for value in [
        "Example E20",
        "Alex Example",
        "Example Street 12",
        "12345 Example Town",
        "REPORTVIN",
    ]:
        assert value in text
    assert "Changed Driver" not in text and "DISCOVERYVIN" not in text
    sheet = load_workbook(io.BytesIO(xlsx_bytes(frozen))).active
    cells = {sheet.cell(1, c).value: sheet.cell(2, c) for c in range(1, sheet.max_column + 1)}
    assert cells["Driver address"].value == v.driver_address
    assert cells["VIN"].value == "REPORTVIN"
    assert cells["Driver name"].data_type == "s"


def test_company_pdf_accepts_reports_without_new_metadata_columns():
    import csv

    frozen = report()
    rows = list(csv.DictReader(io.StringIO(frozen.csv_text)))
    fields = [
        k for k in rows[0] if k not in ("Vehicle model", "Driver name", "Driver address", "VIN")
    ]
    output = io.StringIO()
    writer = csv.DictWriter(output, fieldnames=fields, extrasaction="ignore")
    writer.writeheader()
    writer.writerows(rows)
    frozen.csv_text = output.getvalue()
    text = PdfReader(io.BytesIO(pdf_bytes(frozen))).pages[0].extract_text()
    assert "NW FS 97E" in text and "Tommy" in text


def test_mileage_export_uses_assigned_car_and_german_dates():
    import csv

    v, s, t = sample()
    s.vehicle_id = "car-b"
    s.payload["odometers"] = {"car-a": {"km": "1000"}, "car-b": {"km": "23456"}}
    snapshot, text, digest = freeze_report(v, [s], "2026-10", "Europe/Berlin", [t])
    row = next(csv.DictReader(io.StringIO(text)))
    assert row["Start"] == "01.10.26 12:00"
    assert row["Odometer km"] == "23456"
    frozen = NS(
        id="mileage",
        policy="employer",
        month="2026-10",
        snapshot=snapshot,
        csv_text=text,
        sha256=digest,
    )
    pdf = PdfReader(io.BytesIO(pdf_bytes(frozen))).pages[0].extract_text()
    assert "23456" in pdf and "01.10.26 12:00" in pdf
    sheet = load_workbook(io.BytesIO(xlsx_bytes(frozen))).active
    assert sheet.cell(2, sheet.max_column).value == 23456
    s.vehicle_id = "missing"
    row = next(
        csv.DictReader(io.StringIO(freeze_report(v, [s], "2026-10", "Europe/Berlin", [t])[1]))
    )
    assert row["Odometer km"] == ""


def test_pdf_metadata_has_meaningful_title_and_author():
    reader = PdfReader(io.BytesIO(pdf_bytes(report())))
    assert reader.metadata.title == "Wallbox export - 2026-10"
    assert reader.metadata.author == "WattTally"
