"""Freeze a CSV once; derive both presentation formats from those exact rows."""

import csv
import hashlib
import io
from decimal import ROUND_HALF_UP, Decimal
from xml.sax.saxutils import escape
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.accounting import instant, rate_at

BASE_COLUMNS = ["Session", "Period", "Vehicle", "License plate", "Start", "End", "Energy kWh"]
EMPLOYER_COLUMNS = ["Rate EUR/kWh", "Amount EUR"]
ACTUAL_COLUMNS = ["Grid kWh", "Solar kWh", "Grid cost EUR", "Solar cost EUR", "Amount EUR", "Notes"]
BLOCKING_FLAGS = {
    "partial_start",
    "missing_meter_baseline",
    "meter_gap",
    "meter_reset",
    "state_gap",
}


def fixed(n, places="0.001"):
    return str(Decimal(str(n)).quantize(Decimal(places), rounding=ROUND_HALF_UP))


def safe_text(value):
    text = str(value)
    return (
        "'" + text
        if text.lstrip().startswith(("=", "+", "-", "@")) or text.startswith(("\t", "\r", "\n"))
        else text
    )


def freeze_report(vehicle, sessions, month, timezone, tariffs):
    zone = ZoneInfo(timezone)
    rows = []
    for session in sessions:
        if instant(session.start).astimezone(zone).strftime("%Y-%m") != month:
            continue
        p = session.payload
        energy = Decimal(p["energy_kwh"])
        if energy <= 0 or not session.end:
            continue
        reviewed = getattr(session, "reviewed_energy_kwh", None)
        if session.assignment != "confirmed" or reviewed is None or Decimal(reviewed) != energy:
            continue
        row = {
            "Session": session.id,
            "Period": month,
            "Vehicle": safe_text(vehicle.name),
            "License plate": safe_text(vehicle.plate),
            "Start": instant(session.start).astimezone(zone).isoformat(timespec="seconds"),
            "End": instant(session.end).astimezone(zone).isoformat(timespec="seconds"),
            "Energy kWh": p["energy_kwh"],
        }
        if vehicle.policy == "employer":
            tariff = rate_at(tariffs, instant(session.start))
            if tariff is None:
                raise ValueError("Add a reimbursement rate covering this session.")
            rate = tariff.reimbursement_rate
            row.update(
                {
                    "Rate EUR/kWh": rate,
                    "Amount EUR": fixed(Decimal(p["energy_kwh"]) * Decimal(rate), "0.01"),
                }
            )
        else:
            unknown = Decimal(p["unknown_kwh"])
            if "missing_tariff" in p["flags"]:
                raise ValueError(
                    "Add a home electricity tariff in Connections & cars covering "
                    + row["Start"][:10]
                    + ", then create the report again."
                )
            if unknown > min(Decimal("0.1"), energy * Decimal("0.01")):
                raise ValueError(
                    "Charge on "
                    + row["Start"][:10]
                    + " has "
                    + fixed(unknown)
                    + " kWh without grid/solar data. Actual costs cannot be calculated "
                    "reliably; company reimbursement does not require this split."
                )
            known = Decimal(p["grid_kwh"]) + Decimal(p["solar_kwh"])
            if known <= 0:
                raise ValueError("No grid/solar allocation available for " + row["Start"][:10])
            factor = (known + unknown) / known
            notes = []
            if BLOCKING_FLAGS.intersection(p["flags"]):
                notes.append("Energy checked and accepted by user despite meter history gaps.")
            if unknown > 0:
                notes.append(
                    f"Estimated split and cost for {fixed(unknown)} kWh using this charge's known proportions."
                )
            row.update(
                {
                    "Grid kWh": fixed(Decimal(p["grid_kwh"]) * factor, "0.000001"),
                    "Solar kWh": fixed(Decimal(p["solar_kwh"]) * factor, "0.000001"),
                    "Grid cost EUR": fixed(Decimal(p["grid_cost"]) * factor, "0.000001"),
                    "Solar cost EUR": fixed(Decimal(p["solar_cost"]) * factor, "0.000001"),
                    "Amount EUR": fixed(
                        (Decimal(p["grid_cost"]) + Decimal(p["solar_cost"])) * factor, "0.01"
                    ),
                    "Notes": " ".join(notes),
                }
            )
        rows.append(row)
    if not rows:
        raise ValueError(
            "No confirmed charges for this vehicle and month. In Sessions, check the car and kWh, then choose Confirm charge."
        )
    columns = BASE_COLUMNS + (EMPLOYER_COLUMNS if vehicle.policy == "employer" else ACTUAL_COLUMNS)
    out = io.StringIO(newline="")
    writer = csv.DictWriter(out, fieldnames=columns)
    writer.writeheader()
    writer.writerows(rows)
    text = out.getvalue()
    snapshot = {
        "review_rule": "Only completed charges whose car and energy total were confirmed are included.",
        "timezone": timezone,
        "vehicle": vehicle.name,
        "plate": vehicle.plate,
        "policy": vehicle.policy,
        "month": month,
        "session_count": len(rows),
        "energy_kwh": fixed(sum(Decimal(r["Energy kWh"]) for r in rows)),
        "amount_eur": fixed(sum(Decimal(r["Amount EUR"]) for r in rows), "0.01"),
        "period_rule": "Whole sessions belong to the month in which charging starts.",
        "rounding": "Round each session to cents, then sum.",
    }
    return snapshot, text, hashlib.sha256(text.encode("utf-8")).hexdigest()


def table_from_csv(report):
    return list(csv.reader(io.StringIO(report.csv_text)))


def xlsx_bytes(report):
    wb = Workbook()
    ws = wb.active
    ws.title = "Charging sessions"
    for row in table_from_csv(report):
        ws.append(row)
    for row in ws.iter_rows():
        for cell in row:
            cell.data_type = "s"  # Treat all user-supplied cells as text, never formulas.
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="6B4F3B")
    for column in ws.columns:
        ws.column_dimensions[column[0].column_letter].width = min(
            38, max(14, max(len(str(c.value or "")) for c in column) + 2)
        )
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for row in ws.iter_rows(min_row=2, min_col=7):
        for cell in row:
            if ws.cell(1, cell.column).value == "Notes":
                continue
            cell.value = float(Decimal(cell.value))
            cell.number_format = "0.000" if "kWh" in ws.cell(1, cell.column).value else "0.00"
    summary = wb.create_sheet("Report")
    for key, value in {
        **report.snapshot,
        "report_id": report.id,
        "csv_sha256": report.sha256,
    }.items():
        summary.append([key, str(value)])
        summary.cell(summary.max_row, 2).data_type = "s"
    summary.column_dimensions["A"].width = 22
    summary.column_dimensions["B"].width = 80
    stream = io.BytesIO()
    wb.save(stream)
    return stream.getvalue()


def pdf_bytes(report):
    stream = io.BytesIO()
    doc = SimpleDocTemplate(
        stream, pagesize=landscape(A4), rightMargin=28, leftMargin=28, topMargin=28, bottomMargin=38
    )
    styles = getSampleStyleSheet()
    title = "Charging reimbursement" if report.policy == "employer" else "Charging costs"
    story = [
        Paragraph(title, styles["Title"]),
        Paragraph(
            escape(f"{report.snapshot['vehicle']} · {report.snapshot['plate']} · {report.month}"),
            styles["Heading2"],
        ),
    ]
    source = list(csv.DictReader(io.StringIO(report.csv_text)))
    headers = ["#", "Start", "End", "Energy kWh"]
    headers += (
        ["EUR/kWh", "Amount EUR"]
        if report.policy == "employer"
        else ["Grid kWh", "Solar kWh", "Amount EUR"]
    )
    rows = [headers]
    for i, r in enumerate(source, 1):
        values = [
            str(i),
            r["Start"].replace("T", " ")[:19],
            r["End"].replace("T", " ")[:19],
            fixed(r["Energy kWh"]),
        ]
        values += (
            [r["Rate EUR/kWh"], r["Amount EUR"]]
            if report.policy == "employer"
            else [fixed(r["Grid kWh"]), fixed(r["Solar kWh"]), r["Amount EUR"]]
        )
        rows.append(values)
    table = Table(rows, repeatRows=1, hAlign="LEFT")
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#6b4f3b")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 9),
                ("TOPPADDING", (0, 0), (-1, -1), 9),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.HexColor("#f8f3ee"), colors.white]),
                ("LINEBELOW", (0, 0), (-1, 0), 0.5, colors.HexColor("#d9c5b3")),
            ]
        )
    )
    story.extend(
        [
            Spacer(1, 14),
            table,
            Spacer(1, 18),
            Paragraph(
                f"Total: {report.snapshot['energy_kwh']} kWh · EUR {report.snapshot['amount_eur']}",
                styles["Heading2"],
            ),
            Paragraph(
                escape(
                    f"Times: {report.snapshot['timezone']}. {report.snapshot['period_rule']} {report.snapshot['rounding']}"
                ),
                styles["Normal"],
            ),
        ]
    )
    if report.policy == "actual":
        story.append(
            Paragraph(
                "Solar allocation follows the house-first rule using charger power and grid import. Local supply is labelled solar; battery provenance is not tracked.",
                styles["Normal"],
            )
        )

    if report.policy == "actual":
        for index, row in enumerate(source, 1):
            if row.get("Notes"):
                story.append(Paragraph(escape(f"Charge {index}: {row['Notes']}"), styles["Normal"]))

    def footer(canvas, document):
        canvas.setFont("Helvetica", 7)
        canvas.drawString(28, 20, f"WattTally · {report.id} · CSV SHA256 {report.sha256}")
        canvas.drawRightString(812, 20, str(document.page))

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return stream.getvalue()
